"""Convert complete-documentation.md to .docx"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "documentation" / "complete-project" / "complete-documentation.md"
DOCX_PATH = ROOT / "documentation" / "complete-project" / "complete-documentation.docx"


def set_run_font(run, size=11, bold=False, italic=False, code=False, color=None):
    run.bold = bold
    run.italic = italic
    if code:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    else:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def add_formatted_runs(paragraph, content, base_size=11):
    pattern = re.compile(
        r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))"
    )
    pos = 0
    for m in pattern.finditer(content):
        if m.start() > pos:
            run = paragraph.add_run(content[pos : m.start()])
            set_run_font(run, size=base_size)
        token = m.group(0)
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, code=True)
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, italic=True)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            set_run_font(run, size=base_size, color=RGBColor(0x00, 0x66, 0xCC))
        pos = m.end()
    if pos < len(content):
        run = paragraph.add_run(content[pos:])
        set_run_font(run, size=base_size)


def add_code_block(doc, code_lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.15)
    run = p.add_run("\n".join(code_lines))
    set_run_font(run, code=True)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F4F4")
    shd.set(qn("w:val"), "clear")
    p.paragraph_format.element.get_or_add_pPr().append(shd)


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def is_table_separator(line):
    return bool(
        re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line)
    )


def parse_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        add_formatted_runs(p, h, base_size=10)
        for run in p.runs:
            run.bold = True
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8F4FC")
        shd.set(qn("w:val"), "clear")
        tc_pr.append(shd)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            if c_i >= len(header):
                break
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            add_formatted_runs(cell.paragraphs[0], val, base_size=10)
    doc.add_paragraph()


def main():
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        raw = line
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            add_code_block(doc, buf)
            i += 1
            continue

        if re.match(r"^---+$", stripped) or re.match(r"^\*\*\*+$", stripped):
            add_hr(doc)
            i += 1
            continue

        hm = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if hm:
            level = len(hm.group(1))
            content = hm.group(2).strip()
            style_name = f"Heading {min(level, 9)}"
            p = doc.add_paragraph(style=style_name)
            add_formatted_runs(p, content, base_size=14 if level == 1 else 12)
            i += 1
            continue

        if "|" in stripped and i + 1 < n and is_table_separator(lines[i + 1].strip()):
            header = parse_table_row(stripped)
            i += 2
            rows = []
            while i < n and "|" in lines[i] and not is_table_separator(lines[i].strip()):
                if lines[i].strip():
                    rows.append(parse_table_row(lines[i]))
                i += 1
            width = len(header)
            norm = []
            for r in rows:
                if len(r) < width:
                    r = r + [""] * (width - len(r))
                norm.append(r[:width])
            add_table(doc, header, norm)
            continue

        if stripped.startswith(">"):
            content = re.sub(r"^>\s?", "", stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            add_formatted_runs(p, content, base_size=11)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        um = re.match(r"^(\s*)([-*+])\s+(.*)$", raw)
        if um:
            indent, _, content = um.groups()
            level = len(indent) // 2
            p = doc.add_paragraph(style="List Bullet")
            if level:
                p.paragraph_format.left_indent = Inches(0.25 * level)
            add_formatted_runs(p, content)
            i += 1
            continue

        om = re.match(r"^(\s*)(\d+)\.\s+(.*)$", raw)
        if om:
            indent, _, content = om.groups()
            level = len(indent) // 2
            p = doc.add_paragraph(style="List Number")
            if level:
                p.paragraph_format.left_indent = Inches(0.25 * level)
            add_formatted_runs(p, content)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_formatted_runs(p, stripped)
        i += 1

    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH} ({DOCX_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
